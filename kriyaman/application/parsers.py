import io
import os
import re
from bs4 import BeautifulSoup
import docx
from pypdf import PdfReader
from domain.errors import ValidationError

SUPPORTED_EXTENSIONS = {".txt", ".md", ".html", ".htm", ".pdf", ".docx"}


def sanitize_filename(filename: str) -> str:
    """Sanitize filename to prevent directory traversal or control characters."""
    clean = os.path.basename(filename).strip()
    clean = re.sub(r"[^\w\.\-\s\(\)]", "_", clean)
    clean = re.sub(r"\s+", " ", clean)
    if not clean or clean.startswith("."):
        return f"document_{hash(filename) % 10000}"
    return clean


def detect_mime_type(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext in [".txt"]:
        return "text/plain"
    if ext in [".md"]:
        return "text/markdown"
    if ext in [".html", ".htm"]:
        return "text/html"
    if ext in [".pdf"]:
        return "application/pdf"
    if ext in [".docx"]:
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    return "application/octet-stream"


class DocumentParser:
    """Extract clean plain text from supported document formats."""

    @classmethod
    def parse(cls, filename: str, content: bytes) -> tuple[str, dict[str, str | int]]:
        """Parse raw document bytes into cleaned text and metadata."""
        ext = os.path.splitext(filename.lower())[1]
        if ext not in SUPPORTED_EXTENSIONS:
            raise ValidationError(
                f"Unsupported file format '{ext}'. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        if not content:
            raise ValidationError("Document content is empty.")

        metadata: dict[str, str | int] = {
            "original_filename": filename,
            "extension": ext,
            "byte_size": len(content),
        }

        if ext in [".txt", ".md"]:
            text = cls._parse_text(content)
        elif ext in [".html", ".htm"]:
            text = cls._parse_html(content)
        elif ext == ".pdf":
            text = cls._parse_pdf(content, metadata)
        elif ext == ".docx":
            text = cls._parse_docx(content, metadata)
        else:
            raise ValidationError(f"Handler missing for supported extension: {ext}")

        cleaned_text = cls._clean_text(text)
        if not cleaned_text.strip():
            raise ValidationError(f"No readable text could be extracted from '{filename}'.")

        return cleaned_text, metadata

    @staticmethod
    def _parse_text(content: bytes) -> str:
        for encoding in ["utf-8", "latin-1", "cp1252"]:
            try:
                return content.decode(encoding)
            except UnicodeDecodeError:
                continue
        return content.decode("utf-8", errors="replace")

    @staticmethod
    def _parse_html(content: bytes) -> str:
        soup = BeautifulSoup(content, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header"]):
            tag.decompose()
        return soup.get_text(separator="\n")

    @staticmethod
    def _parse_pdf(content: bytes, metadata: dict[str, str | int]) -> str:
        try:
            reader = PdfReader(io.BytesIO(content))
            pages_text = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    pages_text.append(page_text)
            metadata["page_count"] = len(reader.pages)
            return "\n\n".join(pages_text)
        except Exception as exc:
            raise ValidationError(f"Failed to parse PDF document: {exc}") from exc

    @staticmethod
    def _parse_docx(content: bytes, metadata: dict[str, str | int]) -> str:
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            metadata["paragraph_count"] = len(doc.paragraphs)
            return "\n\n".join(paragraphs)
        except Exception as exc:
            raise ValidationError(f"Failed to parse DOCX document: {exc}") from exc

    @staticmethod
    def _clean_text(text: str) -> str:
        # Normalize whitespace while preserving paragraph breaks
        lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
        cleaned = "\n".join(lines)
        return re.sub(r"\n{3,}", "\n\n", cleaned).strip()
