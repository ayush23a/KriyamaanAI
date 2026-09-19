import io
import docx
from pypdf import PdfWriter
import pytest
from domain.errors import ValidationError
from application.parsers import DocumentParser, sanitize_filename


def test_sanitize_filename():
    assert sanitize_filename("../../etc/passwd") == "passwd"
    assert sanitize_filename("my report (v1).pdf") == "my report (v1).pdf"
    assert sanitize_filename("   ") != ""


def test_parse_txt():
    content = b"Hello world!\nThis is a simple text document."
    text, meta = DocumentParser.parse("sample.txt", content)
    assert "Hello world!" in text
    assert meta["extension"] == ".txt"
    assert meta["byte_size"] == len(content)


def test_parse_markdown():
    content = b"# Title\n\nThis is **bold** text in markdown."
    text, meta = DocumentParser.parse("notes.md", content)
    assert "# Title" in text
    assert "bold" in text


def test_parse_html():
    content = b"""
    <html>
      <head><title>Test</title><style>body { color: red; }</style></head>
      <body>
        <script>alert('bad');</script>
        <header>Header content</header>
        <h1>Main Heading</h1>
        <p>This is paragraph content.</p>
      </body>
    </html>
    """
    text, meta = DocumentParser.parse("index.html", content)
    assert "alert" not in text
    assert "body { color: red; }" not in text
    assert "Main Heading" in text
    assert "This is paragraph content." in text


def test_parse_docx():
    doc = docx.Document()
    doc.add_heading("Docx Heading", 0)
    doc.add_paragraph("First paragraph in Word document.")
    buf = io.BytesIO()
    doc.save(buf)
    content = buf.getvalue()

    text, meta = DocumentParser.parse("test.docx", content)
    assert "Docx Heading" in text
    assert "First paragraph in Word document." in text
    assert meta["extension"] == ".docx"


def test_parse_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buf = io.BytesIO()
    writer.write(buf)
    content = buf.getvalue()

    # A completely blank page has no text, should raise ValidationError
    with pytest.raises(ValidationError):
        DocumentParser.parse("empty.pdf", content)


def test_unsupported_extension():
    with pytest.raises(ValidationError) as exc:
        DocumentParser.parse("file.exe", b"malicious binary")
    assert "Unsupported file format" in str(exc.value)


def test_empty_content():
    with pytest.raises(ValidationError):
        DocumentParser.parse("file.txt", b"")

